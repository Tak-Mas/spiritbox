"""Tests for script sources (Claude / notes) and Gemini 3.x TTS handling.

All tests run without network access; API clients are mocked.
"""

import base64
import io
import wave
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from pptx import Presentation

from spiritbox.config import (
    DEFAULT_SCRIPT_MODEL,
    DEFAULT_TTS_MODEL,
    SpiritBoxConfig,
    load_config,
    save_config,
)
from spiritbox.extract import extract_text_and_notes
from spiritbox.script import (
    _build_prompt,
    _claude_model_options,
    build_script_from_notes,
    generate_master_script,
    parse_master_script,
    slides_without_notes,
)
from spiritbox.voice import (
    _build_tts_request,
    _decode_audio,
    _synthesize_text,
    generate_slide_audio,
    get_wav_duration_ms,
)
from tests.test_spiritbox import _make_pptx


@pytest.fixture
def project_dir(tmp_path):
    for subdir in ["source", "slides", "script", "audio", "output"]:
        (tmp_path / subdir).mkdir()
    return tmp_path


@pytest.fixture
def pdf_path(tmp_path):
    path = tmp_path / "deck.pdf"
    path.write_bytes(b"%PDF-1.4 fake")
    return str(path)


SLIDES = [
    {"slide_number": 1, "title": "Intro", "text": "Intro", "notes": "  こんにちは。\n"},
    {"slide_number": 2, "title": "Body", "text": "Body", "notes": ""},
    {"slide_number": 3, "title": "End", "text": "End", "notes": "以上です。"},
]


def _wav_bytes(pcm: bytes, rate: int) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(rate)
        wf.writeframes(pcm)
    return buf.getvalue()


# ===========================================================================
# config
# ===========================================================================

class TestConfigNewFields:

    def test_defaults(self):
        cfg = SpiritBoxConfig()
        assert cfg.script_source == "generate"
        assert cfg.output_format == "ppsx"
        assert cfg.script_model == DEFAULT_SCRIPT_MODEL == "claude-opus-5-5"
        assert cfg.tts_model == DEFAULT_TTS_MODEL == "gemini-3.8-flash-tts"

    def test_round_trip(self, tmp_path):
        cfg = SpiritBoxConfig(
            script_source="notes_refine",
            output_format="both",
            script_model="gemini-3-pro-preview",
            tts_model="gemini-2.5-pro-preview-tts",
        )
        save_config(cfg, str(tmp_path))
        loaded = load_config(str(tmp_path))
        assert loaded.script_source == "notes_refine"
        assert loaded.output_format == "both"
        assert loaded.script_model == "gemini-3-pro-preview"
        assert loaded.tts_model == "gemini-2.5-pro-preview-tts"

    def test_old_config_gets_new_defaults(self, tmp_path):
        (tmp_path / "config.yaml").write_text("language: en\ncontext: old\n")
        loaded = load_config(str(tmp_path))
        assert loaded.script_source == "generate"
        assert loaded.tts_model == DEFAULT_TTS_MODEL

    @pytest.mark.parametrize("fmt, expected", [("ppsx", False), ("video", True), ("both", True)])
    def test_is_video(self, fmt, expected):
        assert SpiritBoxConfig(output_format=fmt).is_video is expected

    def test_validate_rejects_unknown_values(self):
        with pytest.raises(ValueError, match="script_source"):
            SpiritBoxConfig(script_source="auto").validate()
        with pytest.raises(ValueError, match="output_format"):
            SpiritBoxConfig(output_format="gif").validate()


# ===========================================================================
# extract: title and hidden flag
# ===========================================================================

class TestExtractTitleAndHidden:

    def test_title_from_placeholder(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = "Quarterly  Results"
        path = tmp_path / "t.pptx"
        prs.save(str(path))
        data = extract_text_and_notes(str(path))
        assert data[0]["title"] == "Quarterly Results"
        assert data[0]["hidden"] is False

    def test_title_falls_back_to_first_text_line(self, tmp_path):
        path = tmp_path / "t.pptx"
        _make_pptx(path, [{"text": "First line\nSecond line"}, {}])
        data = extract_text_and_notes(str(path))
        assert data[0]["title"] == "First line"
        assert data[1]["title"] == ""

    def test_hidden_flag(self, tmp_path):
        path = tmp_path / "t.pptx"
        _make_pptx(path, [{"text": "a"}, {"text": "b"}])
        prs = Presentation(str(path))
        prs.slides[1]._element.set("show", "0")
        prs.save(str(path))
        data = extract_text_and_notes(str(path))
        assert [s["hidden"] for s in data] == [False, True]


# ===========================================================================
# script: notes as the script source
# ===========================================================================

class TestNotesScriptSource:

    def test_build_script_from_notes(self):
        text = build_script_from_notes(SLIDES)
        assert "[---SLIDE_01---]\nこんにちは。" in text
        assert "[---SLIDE_02---]" in text
        assert "[---SLIDE_03---]\n以上です。" in text

    def test_slides_without_notes(self):
        assert slides_without_notes(SLIDES) == [2]

    def test_notes_mode_calls_no_api(self, project_dir, pdf_path):
        cfg = SpiritBoxConfig(script_source="notes")
        with patch("spiritbox.script._generate_with_claude") as claude, \
                patch("spiritbox.script._generate_with_gemini") as gemini:
            generate_master_script(str(project_dir), cfg, SLIDES, pdf_path)
        claude.assert_not_called()
        gemini.assert_not_called()

        scripts = parse_master_script(str(project_dir))
        # Slide 2 has no notes → no narration
        assert scripts == {1: "こんにちは。", 3: "以上です。"}

    def test_existing_script_is_kept(self, project_dir, pdf_path):
        (project_dir / "script" / "master_script.md").write_text("[---SLIDE_01---]\nedited")
        cfg = SpiritBoxConfig(script_source="notes")
        text = generate_master_script(str(project_dir), cfg, SLIDES, pdf_path)
        assert text == "[---SLIDE_01---]\nedited"


class TestPromptModes:

    def test_generate_treats_notes_as_private(self):
        prompt = _build_prompt(3, SpiritBoxConfig(script_source="generate"))
        assert "private notes" in prompt
        assert "must not be read aloud" in prompt

    def test_notes_refine_keeps_author_content(self):
        prompt = _build_prompt(3, SpiritBoxConfig(script_source="notes_refine"))
        assert "draft narration" in prompt
        assert "Keep the author's content" in prompt

    def test_video_prompt(self):
        prompt = _build_prompt(3, SpiritBoxConfig(output_format="video"))
        assert "recorded video" in prompt
        assert "Do not use filler words" in prompt

    def test_ppsx_prompt_keeps_live_style(self):
        prompt = _build_prompt(3, SpiritBoxConfig(output_format="ppsx"))
        assert "live presentation" in prompt

    def test_plain_text_rule_and_last_marker(self):
        prompt = _build_prompt(12, SpiritBoxConfig())
        assert "Do not use Markdown" in prompt
        assert "[---SLIDE_12---]" in prompt


# ===========================================================================
# script: Claude backend
# ===========================================================================

def _mock_anthropic(text="[---SLIDE_01---]\nHello", stop_reason="end_turn"):
    """Patch anthropic.Anthropic; return (patcher, stream kwargs recorder)."""
    message = SimpleNamespace(
        stop_reason=stop_reason,
        stop_details=SimpleNamespace(category="cyber") if stop_reason == "refusal" else None,
        content=[SimpleNamespace(type="thinking", thinking=""),
                 SimpleNamespace(type="text", text=text)],
    )
    stream = MagicMock()
    stream.__enter__.return_value.get_final_message.return_value = message
    client = MagicMock()
    client.beta.messages.stream.return_value = stream
    return patch("anthropic.Anthropic", return_value=client), client


class TestClaudeBackend:

    def test_request_shape_and_saved_script(self, project_dir, pdf_path):
        patcher, client = _mock_anthropic("Here is the script.\n[---SLIDE_01---]\nHello")
        with patcher:
            text = generate_master_script(
                str(project_dir), SpiritBoxConfig(), SLIDES, pdf_path
            )

        kwargs = client.beta.messages.stream.call_args.kwargs
        assert kwargs["model"] == "claude-opus-5-5"
        assert kwargs["fallbacks"] == "default"
        assert kwargs["betas"] == ["server-side-fallback-2026-07-01"]
        assert kwargs["output_config"] == {"effort": "high"}

        content = kwargs["messages"][0]["content"]
        assert content[0]["type"] == "document"
        assert base64.b64decode(content[0]["source"]["data"]) == b"%PDF-1.4 fake"
        assert "Speaker Notes: こんにちは。" in content[1]["text"]
        assert "[---SLIDE_03---]" in content[2]["text"]

        # Preamble before the first marker is dropped
        assert text == "[---SLIDE_01---]\nHello"
        assert parse_master_script(str(project_dir)) == {1: "Hello"}

    def test_refusal_raises(self, project_dir, pdf_path):
        patcher, _ = _mock_anthropic(stop_reason="refusal")
        with patcher, pytest.raises(RuntimeError, match="declined"):
            generate_master_script(str(project_dir), SpiritBoxConfig(), SLIDES, pdf_path)
        assert not (project_dir / "script" / "master_script.md").exists()

    def test_truncation_raises(self, project_dir, pdf_path):
        patcher, _ = _mock_anthropic(stop_reason="max_tokens")
        with patcher, pytest.raises(RuntimeError, match="cut off"):
            generate_master_script(str(project_dir), SpiritBoxConfig(), SLIDES, pdf_path)

    def test_large_pdf_rejected(self, project_dir, pdf_path):
        with patch("spiritbox.script._CLAUDE_MAX_PDF_BYTES", 4), \
                pytest.raises(RuntimeError, match="MB"):
            generate_master_script(str(project_dir), SpiritBoxConfig(), SLIDES, pdf_path)

    def test_notes_refine_uses_claude(self, project_dir, pdf_path):
        patcher, client = _mock_anthropic()
        cfg = SpiritBoxConfig(script_source="notes_refine")
        with patcher:
            generate_master_script(str(project_dir), cfg, SLIDES, pdf_path)
        prompt = client.beta.messages.stream.call_args.kwargs["messages"][0]["content"][2]["text"]
        assert "draft narration" in prompt

    def test_older_models_get_no_new_options(self):
        assert _claude_model_options("claude-haiku-4-5") == {}
        assert "fallbacks" in _claude_model_options("claude-sonnet-5-5")

    def test_gemini_model_dispatch(self, project_dir, pdf_path):
        cfg = SpiritBoxConfig(script_model="gemini-3-pro-preview")
        with patch("spiritbox.script._generate_with_gemini",
                   return_value="[---SLIDE_01---]\nG") as gemini:
            generate_master_script(str(project_dir), cfg, SLIDES, pdf_path)
        gemini.assert_called_once()

    def test_unknown_model_raises(self, project_dir, pdf_path):
        cfg = SpiritBoxConfig(script_model="gpt-x")
        with pytest.raises(ValueError, match="script_model"):
            generate_master_script(str(project_dir), cfg, SLIDES, pdf_path)


# ===========================================================================
# voice: Gemini 3.x TTS
# ===========================================================================

class TestTtsRequest:

    def test_gemini3_uses_voice_name_and_style(self):
        cfg = SpiritBoxConfig(language="ja", voice_style="落ち着いた声")
        contents, gen_cfg = _build_tts_request("こんにちは", cfg)
        assert gen_cfg.response_modalities == ["AUDIO"]
        assert gen_cfg.speech_config.voice_config.voice == "Aoede"
        part = contents[0].parts[0]
        assert part.text == "こんにちは"
        assert part.speech_metadata.style == "落ち着いた声"

    def test_gemini3_without_style(self):
        contents, _ = _build_tts_request("hi", SpiritBoxConfig(language="en"))
        assert contents[0].parts[0].speech_metadata is None

    def test_gemini2_keeps_prebuilt_voice(self):
        cfg = SpiritBoxConfig(tts_model="gemini-2.5-pro-preview-tts", voice="Puck")
        contents, gen_cfg = _build_tts_request("hi", cfg)
        assert contents == "hi"
        voice_config = gen_cfg.speech_config.voice_config
        assert voice_config.prebuilt_voice_config.voice_name == "Puck"
        assert voice_config.voice is None


class TestDecodeAudio:

    def test_wav_container(self):
        pcm = b"\x01\x00" * 100
        data, rate = _decode_audio(_wav_bytes(pcm, 24000), "audio/wav")
        assert data == pcm
        assert rate == 24000

    def test_raw_pcm_rate_from_mime(self):
        data, rate = _decode_audio(b"\x00\x00" * 10, "audio/L16;codec=pcm;rate=16000")
        assert data == b"\x00\x00" * 10
        assert rate == 16000

    def test_raw_pcm_default_rate(self):
        _, rate = _decode_audio(b"\x00\x00", None)
        assert rate == 24000

    def test_stereo_wav_rejected(self):
        buf = io.BytesIO()
        with wave.open(buf, "wb") as wf:
            wf.setnchannels(2)
            wf.setsampwidth(2)
            wf.setframerate(24000)
            wf.writeframes(b"\x00" * 8)
        with pytest.raises(RuntimeError, match="channel"):
            _decode_audio(buf.getvalue(), "audio/wav")


def _tts_response(data=None, mime_type="audio/wav"):
    if data is None:
        return SimpleNamespace(candidates=[SimpleNamespace(
            content=SimpleNamespace(parts=[]), finish_reason="OTHER")])
    part = SimpleNamespace(inline_data=SimpleNamespace(data=data, mime_type=mime_type))
    return SimpleNamespace(candidates=[SimpleNamespace(
        content=SimpleNamespace(parts=[part]), finish_reason="STOP")])


class TestSynthesize:

    def test_retries_empty_audio(self):
        pcm = b"\x02\x00" * 50
        client = MagicMock()
        client.models.generate_content.side_effect = [
            _tts_response(None),
            _tts_response(_wav_bytes(pcm, 24000)),
        ]
        with patch("spiritbox.voice.time.sleep"):
            data, rate = _synthesize_text(client, "hi", SpiritBoxConfig())
        assert (data, rate) == (pcm, 24000)
        assert client.models.generate_content.call_count == 2
        assert client.models.generate_content.call_args.kwargs["model"] == "gemini-3.8-flash-tts"

    def test_gives_up_after_max_attempts(self):
        client = MagicMock()
        client.models.generate_content.return_value = _tts_response(None)
        with patch("spiritbox.voice.time.sleep"), \
                pytest.raises(RuntimeError, match="no audio"):
            _synthesize_text(client, "hi", SpiritBoxConfig())

    def test_slide_audio_keeps_sample_rate(self, tmp_path):
        client = MagicMock()
        client.models.generate_content.return_value = _tts_response(
            _wav_bytes(b"\x00\x00" * 4800, 48000))
        out = tmp_path / "s.wav"
        duration = generate_slide_audio(client, "hi", SpiritBoxConfig(), str(out))
        assert duration == 100  # 4800 samples at 48 kHz
        assert get_wav_duration_ms(str(out)) == 100
