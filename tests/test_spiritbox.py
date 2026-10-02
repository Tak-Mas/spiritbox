"""Comprehensive unit tests for the SpiritBox project.

All tests run WITHOUT Google Cloud API access.
External dependencies (Gemini, TTS, LibreOffice) are mocked.
"""

import io
import struct
import wave
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

# ---------------------------------------------------------------------------
# Module imports under test
# ---------------------------------------------------------------------------
from spiritbox.config import (
    SpiritBoxConfig,
    get_source_pptx,
    init_project,
    load_config,
    save_config,
)
from spiritbox.extract import extract_all, extract_text_and_notes
from spiritbox.script import (
    _build_duration_instruction,
    _build_prompt,
    parse_master_script,
)
from spiritbox.voice import (
    _TTS_MAX_CHARS,
    _save_pcm_as_wav,
    _split_text_for_tts,
    generate_all_audio,
    get_wav_duration_ms,
)
from spiritbox.assemble import (
    _convert_to_ppsx,
    _create_poster_image,
    _set_slide_advance_time,
)


# ===========================================================================
# Helpers / Fixtures
# ===========================================================================

def _make_pptx(path, slides_content=None):
    """Create a minimal PPTX at *path* with optional slide text/notes.

    slides_content: list of dicts with keys 'text' and/or 'notes'.
    If None, creates a single blank slide.
    """
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank layout

    if slides_content is None:
        slides_content = [{}]

    for slide_info in slides_content:
        slide = prs.slides.add_slide(blank_layout)

        text = slide_info.get("text", "")
        if text:
            from pptx.util import Inches
            txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            txBox.text_frame.text = text

        notes = slide_info.get("notes", "")
        if notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = notes

    prs.save(str(path))
    return path


def _make_pcm_data(num_samples=4800, sample_rate=24000):
    """Generate silent PCM data (LINEAR16, mono).

    Default: 4800 samples at 24000 Hz = 200ms of audio.
    """
    return b"\x00\x00" * num_samples


@pytest.fixture
def default_config():
    return SpiritBoxConfig()


@pytest.fixture
def ja_config():
    return SpiritBoxConfig(language="ja")


@pytest.fixture
def en_config():
    return SpiritBoxConfig(language="en")


@pytest.fixture
def project_dir(tmp_path):
    """Create a minimal project directory structure."""
    for subdir in ["source", "slides", "script", "audio", "output"]:
        (tmp_path / subdir).mkdir()
    return tmp_path


@pytest.fixture
def project_with_pptx(project_dir):
    """Project directory with a single source PPTX."""
    pptx_path = project_dir / "source" / "test_deck.pptx"
    _make_pptx(pptx_path, [
        {"text": "Slide 1 Title", "notes": "Speaker notes 1"},
        {"text": "Slide 2 Content", "notes": "Speaker notes 2"},
    ])
    return project_dir


# ===========================================================================
# config.py tests
# ===========================================================================

class TestSpiritBoxConfig:
    """Tests for the SpiritBoxConfig dataclass."""

    def test_defaults(self):
        cfg = SpiritBoxConfig()
        assert cfg.language == "ja"
        assert cfg.duration_minutes == 0
        assert cfg.voice == ""
        assert cfg.voice_style == ""
        assert cfg.source_filename == ""
        assert cfg.context == ""

    def test_language_code_ja(self):
        cfg = SpiritBoxConfig(language="ja")
        assert cfg.language_code == "ja-JP"

    def test_language_code_en(self):
        cfg = SpiritBoxConfig(language="en")
        assert cfg.language_code == "en-US"

    def test_language_code_invalid_raises(self):
        cfg = SpiritBoxConfig(language="fr")
        with pytest.raises(ValueError, match="Unsupported language.*fr"):
            _ = cfg.language_code

    def test_language_code_empty_raises(self):
        cfg = SpiritBoxConfig(language="")
        with pytest.raises(ValueError):
            _ = cfg.language_code

    def test_effective_voice_default_ja(self):
        cfg = SpiritBoxConfig(language="ja")
        assert cfg.effective_voice == "Aoede"

    def test_effective_voice_default_en(self):
        cfg = SpiritBoxConfig(language="en")
        assert cfg.effective_voice == "Kore"

    def test_effective_voice_custom_override(self):
        cfg = SpiritBoxConfig(language="ja", voice="CustomVoice")
        assert cfg.effective_voice == "CustomVoice"

    def test_effective_voice_unknown_language_fallback(self):
        """When language is not in the default map, fallback is Kore."""
        cfg = SpiritBoxConfig(language="de", voice="")
        assert cfg.effective_voice == "Kore"


class TestSaveLoadConfig:
    """Tests for save_config and load_config round-trip."""

    def test_round_trip_defaults(self, tmp_path):
        cfg = SpiritBoxConfig()
        save_config(cfg, str(tmp_path))
        loaded = load_config(str(tmp_path))
        assert loaded.language == cfg.language
        assert loaded.duration_minutes == cfg.duration_minutes
        assert loaded.voice == cfg.voice
        assert loaded.voice_style == cfg.voice_style
        assert loaded.source_filename == cfg.source_filename
        assert loaded.context == cfg.context

    def test_round_trip_custom_values(self, tmp_path):
        cfg = SpiritBoxConfig(
            language="en",
            duration_minutes=10,
            voice="MyVoice",
            voice_style="cheerful",
            source_filename="deck.pptx",
            context="A talk about testing.",
        )
        save_config(cfg, str(tmp_path))
        loaded = load_config(str(tmp_path))
        assert loaded.language == "en"
        assert loaded.duration_minutes == 10
        assert loaded.voice == "MyVoice"
        assert loaded.voice_style == "cheerful"
        assert loaded.source_filename == "deck.pptx"
        assert loaded.context == "A talk about testing."

    def test_save_config_stores_raw_voice_not_effective(self, tmp_path):
        """save_config must store the raw voice (''), not effective_voice ('Aoede')."""
        cfg = SpiritBoxConfig(language="ja", voice="")
        assert cfg.effective_voice == "Aoede"  # confirm effective differs
        save_config(cfg, str(tmp_path))

        config_path = tmp_path / "config.yaml"
        with open(config_path) as f:
            data = yaml.safe_load(f)
        assert data["voice"] == ""

    def test_load_config_missing_file_returns_defaults(self, tmp_path):
        loaded = load_config(str(tmp_path))
        assert loaded.language == "ja"
        assert loaded.duration_minutes == 0

    def test_load_config_ignores_unknown_fields(self, tmp_path):
        config_path = tmp_path / "config.yaml"
        with open(config_path, "w") as f:
            yaml.dump({"language": "en", "unknown_field": 42}, f)
        loaded = load_config(str(tmp_path))
        assert loaded.language == "en"

    def test_load_config_empty_yaml(self, tmp_path):
        config_path = tmp_path / "config.yaml"
        config_path.write_text("")
        loaded = load_config(str(tmp_path))
        assert loaded.language == "ja"

    def test_save_config_creates_parent_dirs(self, tmp_path):
        nested = tmp_path / "a" / "b" / "c"
        cfg = SpiritBoxConfig()
        save_config(cfg, str(nested))
        assert (nested / "config.yaml").exists()


class TestInitProject:
    """Tests for init_project."""

    def test_creates_directory_structure(self, tmp_path):
        pptx_path = tmp_path / "input.pptx"
        _make_pptx(pptx_path)

        proj = tmp_path / "my_project"
        result = init_project(str(pptx_path), str(proj))

        assert result == str(proj)
        for subdir in ["source", "slides", "script", "audio", "output"]:
            assert (proj / subdir).is_dir()

    def test_copies_pptx_to_source(self, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        _make_pptx(pptx_path)

        proj = tmp_path / "project"
        init_project(str(pptx_path), str(proj))

        copied = proj / "source" / "deck.pptx"
        assert copied.exists()
        assert copied.stat().st_size == pptx_path.stat().st_size

    def test_does_not_overwrite_existing_source(self, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        _make_pptx(pptx_path)

        proj = tmp_path / "project"
        init_project(str(pptx_path), str(proj))

        dest = proj / "source" / "deck.pptx"
        first_mtime = dest.stat().st_mtime

        # Modify the original to make it different
        _make_pptx(pptx_path, [{"text": "Modified"}])
        init_project(str(pptx_path), str(proj))

        # The source copy should NOT have been overwritten
        assert dest.stat().st_mtime == first_mtime

    def test_default_project_dir(self, tmp_path):
        pptx_path = tmp_path / "my_talk.pptx"
        _make_pptx(pptx_path)

        result = init_project(str(pptx_path))
        expected = str(tmp_path / "my_talk_spiritbox")
        assert result == expected
        assert Path(result).is_dir()

    def test_nonexistent_pptx_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="PPTX not found"):
            init_project(str(tmp_path / "nonexistent.pptx"))


class TestGetSourcePptx:
    """Tests for get_source_pptx."""

    def test_with_source_filename_set(self, project_dir):
        pptx_path = project_dir / "source" / "my_deck.pptx"
        _make_pptx(pptx_path)
        cfg = SpiritBoxConfig(source_filename="my_deck.pptx")
        result = get_source_pptx(str(project_dir), cfg)
        assert result == pptx_path

    def test_source_filename_not_found_falls_back_to_glob(self, project_dir):
        pptx_path = project_dir / "source" / "only_deck.pptx"
        _make_pptx(pptx_path)
        cfg = SpiritBoxConfig(source_filename="nonexistent.pptx")
        result = get_source_pptx(str(project_dir), cfg)
        assert result == pptx_path

    def test_fallback_glob_single_file(self, project_dir):
        pptx_path = project_dir / "source" / "only.pptx"
        _make_pptx(pptx_path)
        result = get_source_pptx(str(project_dir))
        assert result == pptx_path

    def test_fallback_glob_multiple_files_raises(self, project_dir):
        _make_pptx(project_dir / "source" / "a.pptx")
        _make_pptx(project_dir / "source" / "b.pptx")
        with pytest.raises(RuntimeError, match="Multiple PPTX"):
            get_source_pptx(str(project_dir))

    def test_fallback_glob_no_files_raises(self, project_dir):
        with pytest.raises(FileNotFoundError, match="No PPTX found"):
            get_source_pptx(str(project_dir))

    def test_config_none_uses_glob(self, project_dir):
        pptx_path = project_dir / "source" / "deck.pptx"
        _make_pptx(pptx_path)
        result = get_source_pptx(str(project_dir), None)
        assert result == pptx_path


# ===========================================================================
# extract.py tests
# ===========================================================================

class TestExtractTextAndNotes:
    """Tests for extract_text_and_notes."""

    def test_basic_extraction(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [
            {"text": "Hello World", "notes": "Say hello"},
            {"text": "Second Slide", "notes": "More notes"},
        ])
        result = extract_text_and_notes(str(pptx_path))

        assert len(result) == 2
        assert result[0]["slide_number"] == 1
        assert "Hello World" in result[0]["text"]
        assert result[0]["notes"] == "Say hello"
        assert result[1]["slide_number"] == 2
        assert "Second Slide" in result[1]["text"]
        assert result[1]["notes"] == "More notes"

    def test_slide_without_notes(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"text": "No notes here"}])
        result = extract_text_and_notes(str(pptx_path))

        assert len(result) == 1
        assert result[0]["notes"] == ""

    def test_slide_without_text(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"notes": "Only notes"}])
        result = extract_text_and_notes(str(pptx_path))

        assert len(result) == 1
        assert result[0]["text"] == ""
        assert result[0]["notes"] == "Only notes"

    def test_blank_slide(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{}])
        result = extract_text_and_notes(str(pptx_path))

        assert len(result) == 1
        assert result[0]["text"] == ""

    def test_multiple_slides_order(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        slides = [{"text": f"Slide {i}"} for i in range(1, 6)]
        _make_pptx(pptx_path, slides)
        result = extract_text_and_notes(str(pptx_path))

        assert len(result) == 5
        for i, item in enumerate(result, 1):
            assert item["slide_number"] == i


class TestExtractAllResumability:
    """Tests for extract_all resumability logic."""

    def test_skips_pdf_conversion_when_pdf_exists(self, project_with_pptx):
        project = project_with_pptx
        slides_dir = project / "slides"

        # Pre-create the PDF file that extract_all would look for
        (slides_dir / "test_deck.pdf").write_bytes(b"fake_pdf_data")

        with patch("spiritbox.extract.convert_to_pdf") as mock_convert:
            config = SpiritBoxConfig(source_filename="test_deck.pptx")
            result = extract_all(str(project), config)

            # convert_to_pdf should NOT have been called
            mock_convert.assert_not_called()
            assert result["slide_count"] == 2
            assert len(result["slides_data"]) == 2
            assert result["pdf_path"] == str(slides_dir / "test_deck.pdf")

    def test_calls_pdf_conversion_when_pdf_missing(self, project_with_pptx):
        project = project_with_pptx
        slides_dir = project / "slides"

        with patch("spiritbox.extract.convert_to_pdf") as mock_convert:
            # Simulate convert_to_pdf creating the PDF
            def fake_convert(pptx_path, output_dir):
                (Path(output_dir) / "test_deck.pdf").write_bytes(b"fake_pdf")
                return str(Path(output_dir) / "test_deck.pdf")
            mock_convert.side_effect = fake_convert

            config = SpiritBoxConfig(source_filename="test_deck.pptx")
            result = extract_all(str(project), config)

            mock_convert.assert_called_once()
            assert result["slide_count"] == 2
            assert result["pdf_path"] == str(slides_dir / "test_deck.pdf")

    def test_returns_correct_slides_data_structure(self, project_with_pptx):
        project = project_with_pptx
        slides_dir = project / "slides"

        # Pre-create the PDF so conversion is skipped
        (slides_dir / "test_deck.pdf").write_bytes(b"fake_pdf_data")

        with patch("spiritbox.extract.convert_to_pdf") as mock_convert:
            config = SpiritBoxConfig(source_filename="test_deck.pptx")
            result = extract_all(str(project), config)

            mock_convert.assert_not_called()
            assert result["slide_count"] == 2
            assert len(result["slides_data"]) == 2
            assert result["slides_data"][0]["slide_number"] == 1
            assert result["slides_data"][1]["slide_number"] == 2


# ===========================================================================
# script.py tests
# ===========================================================================

class TestBuildPrompt:
    """Tests for _build_prompt."""

    def test_contains_language_ja(self):
        cfg = SpiritBoxConfig(language="ja", context="About AI")
        prompt = _build_prompt(5, cfg)
        assert "Japanese" in prompt
        assert "5 slides" in prompt

    def test_contains_language_en(self):
        cfg = SpiritBoxConfig(language="en")
        prompt = _build_prompt(3, cfg)
        assert "English" in prompt
        assert "3 slides" in prompt

    def test_contains_context(self):
        cfg = SpiritBoxConfig(context="This is a test context about ML.")
        prompt = _build_prompt(1, cfg)
        assert "This is a test context about ML." in prompt

    def test_contains_slide_markers_format(self):
        cfg = SpiritBoxConfig()
        prompt = _build_prompt(2, cfg)
        assert "[---SLIDE_01---]" in prompt
        assert "[---SLIDE_02---]" in prompt

    def test_contains_spoken_language_instruction(self):
        cfg = SpiritBoxConfig()
        prompt = _build_prompt(1, cfg)
        assert "spoken language" in prompt

    def test_contains_transition_instruction(self):
        cfg = SpiritBoxConfig()
        prompt = _build_prompt(1, cfg)
        assert "transition" in prompt.lower()


class TestBuildDurationInstruction:
    """Tests for _build_duration_instruction."""

    def test_no_constraint_returns_empty(self):
        cfg = SpiritBoxConfig(duration_minutes=0)
        result = _build_duration_instruction(5, cfg)
        assert result == ""

    def test_negative_returns_empty(self):
        cfg = SpiritBoxConfig(duration_minutes=-1)
        result = _build_duration_instruction(5, cfg)
        assert result == ""

    def test_with_duration_ja(self):
        cfg = SpiritBoxConfig(language="ja", duration_minutes=10)
        result = _build_duration_instruction(5, cfg)
        assert "10 minutes" in result
        assert "5" in result  # slide count
        assert "120" in result  # 10*60//5 = 120 seconds avg
        assert "Japanese speaking rate" in result

    def test_with_duration_en(self):
        cfg = SpiritBoxConfig(language="en", duration_minutes=15)
        result = _build_duration_instruction(3, cfg)
        assert "15 minutes" in result
        assert "English speaking rate" in result
        assert "300" in result  # 15*60//3 = 300 seconds avg


class TestParseMasterScript:
    """Tests for parse_master_script."""

    def test_well_formed_markers(self, project_dir):
        script_text = (
            "[---SLIDE_01---]\n"
            "Hello, welcome to this presentation.\n\n"
            "[---SLIDE_02---]\n"
            "Now let us discuss the main topic.\n\n"
            "[---SLIDE_03---]\n"
            "Thank you for listening.\n"
        )
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))

        assert len(result) == 3
        assert 1 in result
        assert 2 in result
        assert 3 in result
        assert "welcome" in result[1]
        assert "main topic" in result[2]
        assert "Thank you" in result[3]

    def test_no_markers_returns_empty(self, project_dir):
        script_text = "This is just some text with no markers at all."
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))
        assert result == {}

    def test_malformed_markers_returns_empty(self, project_dir):
        script_text = "[SLIDE_01]\nSome text\n[SLIDE 02]\nMore text"
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))
        assert result == {}

    def test_empty_script_between_markers(self, project_dir):
        script_text = (
            "[---SLIDE_01---]\n"
            "Some content here.\n"
            "[---SLIDE_02---]\n"
            "\n"
            "[---SLIDE_03---]\n"
            "Final slide content.\n"
        )
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))

        # Slide 2 has empty text (stripped), so it should NOT be in results
        assert 1 in result
        assert 2 not in result
        assert 3 in result

    def test_marker_numbering_preserved(self, project_dir):
        script_text = (
            "[---SLIDE_05---]\nFifth slide.\n"
            "[---SLIDE_10---]\nTenth slide.\n"
        )
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))
        assert 5 in result
        assert 10 in result
        assert "Fifth" in result[5]
        assert "Tenth" in result[10]

    def test_leading_text_before_first_marker_ignored(self, project_dir):
        script_text = (
            "Some preamble text that should be ignored.\n\n"
            "[---SLIDE_01---]\n"
            "Actual content.\n"
        )
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text(script_text, encoding="utf-8")

        result = parse_master_script(str(project_dir))
        assert len(result) == 1
        assert "Actual content" in result[1]
        # Preamble should not appear in any value
        for v in result.values():
            assert "preamble" not in v.lower()


# ===========================================================================
# voice.py tests
# ===========================================================================

class TestSplitTextForTts:
    """Tests for _split_text_for_tts."""

    def test_short_text_no_split(self):
        text = "Hello, this is a short text."
        result = _split_text_for_tts(text)
        assert result == [text]

    def test_empty_text(self):
        result = _split_text_for_tts("")
        assert result == [""]

    def test_long_japanese_splits_on_sentence_boundary(self):
        # Create text with multiple sentences that together exceed 4000 characters.
        sentence = "これはテスト文です。"  # 10 characters
        # Need enough sentences to exceed 4000 characters
        text = sentence * 450
        assert len(text) > 4000

        result = _split_text_for_tts(text)
        assert len(result) > 1

        # Each chunk should be within the character limit
        for chunk in result:
            assert len(chunk) <= 4000

        # Reassembled text should equal original
        assert "".join(result) == text

    def test_long_english_splits_on_sentence_boundary(self):
        sentence = "This is a test sentence. "
        text = sentence * 200
        assert len(text) > 4000

        result = _split_text_for_tts(text)
        assert len(result) > 1

        for chunk in result:
            assert len(chunk) <= 4000

    def test_single_sentence_exceeding_limit_force_split(self):
        # One huge sentence with no boundary markers
        text = "a" * 5000  # 5000 chars, well over 4000
        result = _split_text_for_tts(text)
        assert len(result) > 1

        for chunk in result:
            assert len(chunk) <= 4000

        assert "".join(result) == text

    def test_japanese_force_split_multibyte(self):
        # A single long Japanese "sentence" without any period
        text = "あ" * 5000  # 5000 characters, no periods
        assert len(text) > 4000

        result = _split_text_for_tts(text)
        assert len(result) > 1

        for chunk in result:
            assert len(chunk) <= 4000

        assert "".join(result) == text

    def test_text_exactly_at_limit(self):
        # Text exactly at the character limit should NOT be split
        text = "a" * _TTS_MAX_CHARS
        result = _split_text_for_tts(text)
        assert result == [text]

    def test_custom_max_chars(self):
        text = "Hello. World. Foo. Bar."
        result = _split_text_for_tts(text, max_chars=10)
        assert len(result) > 1
        for chunk in result:
            assert len(chunk) <= 10


class TestSavePcmAsWavAndDuration:
    """Tests for _save_pcm_as_wav and get_wav_duration_ms round-trip."""

    def test_round_trip_basic(self, tmp_path):
        wav_path = str(tmp_path / "test.wav")
        num_samples = 24000  # 1 second at 24kHz
        pcm = _make_pcm_data(num_samples=num_samples, sample_rate=24000)
        _save_pcm_as_wav(pcm, wav_path, sample_rate=24000)

        duration = get_wav_duration_ms(wav_path)
        assert duration == 1000  # 1 second

    def test_round_trip_half_second(self, tmp_path):
        wav_path = str(tmp_path / "test.wav")
        num_samples = 12000  # 0.5 seconds at 24kHz
        pcm = _make_pcm_data(num_samples=num_samples)
        _save_pcm_as_wav(pcm, wav_path, sample_rate=24000)

        duration = get_wav_duration_ms(wav_path)
        assert duration == 500

    def test_round_trip_200ms(self, tmp_path):
        wav_path = str(tmp_path / "test.wav")
        num_samples = 4800  # 200ms at 24kHz
        pcm = _make_pcm_data(num_samples=num_samples)
        _save_pcm_as_wav(pcm, wav_path, sample_rate=24000)

        duration = get_wav_duration_ms(wav_path)
        assert duration == 200

    def test_wav_file_is_valid(self, tmp_path):
        wav_path = str(tmp_path / "test.wav")
        pcm = _make_pcm_data(num_samples=100)
        _save_pcm_as_wav(pcm, wav_path)

        with wave.open(wav_path, "rb") as wf:
            assert wf.getnchannels() == 1
            assert wf.getsampwidth() == 2
            assert wf.getframerate() == 24000

    def test_empty_pcm(self, tmp_path):
        wav_path = str(tmp_path / "empty.wav")
        _save_pcm_as_wav(b"", wav_path)
        duration = get_wav_duration_ms(wav_path)
        assert duration == 0

    def test_custom_sample_rate(self, tmp_path):
        wav_path = str(tmp_path / "test.wav")
        num_samples = 44100  # 1 second at 44.1kHz
        pcm = _make_pcm_data(num_samples=num_samples)
        _save_pcm_as_wav(pcm, wav_path, sample_rate=44100)

        duration = get_wav_duration_ms(wav_path)
        assert duration == 1000


class TestGenerateAllAudioErrors:
    """Tests for generate_all_audio error handling."""

    def test_raises_on_empty_markers(self, project_dir):
        """generate_all_audio raises RuntimeError when parse_master_script returns {}."""
        # Create an empty master_script.md (no markers)
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text("No markers here.", encoding="utf-8")

        cfg = SpiritBoxConfig()
        with pytest.raises(RuntimeError, match="No slide markers found"):
            generate_all_audio(str(project_dir), cfg)

    def test_raises_on_empty_markers_mocked(self, project_dir):
        """Same test using mock for clarity."""
        script_path = project_dir / "script" / "master_script.md"
        script_path.write_text("dummy", encoding="utf-8")

        with patch("spiritbox.voice.parse_master_script", return_value={}):
            cfg = SpiritBoxConfig()
            with pytest.raises(RuntimeError, match="No slide markers found"):
                generate_all_audio(str(project_dir), cfg)


# ===========================================================================
# assemble.py tests
# ===========================================================================

class TestCreatePosterImage:
    """Tests for _create_poster_image."""

    def test_returns_valid_png_bytes(self):
        data = _create_poster_image()
        assert isinstance(data, bytes)
        assert len(data) > 0

        # Verify it is valid PNG by opening with PIL
        img = Image.open(io.BytesIO(data))
        assert img.format == "PNG"
        assert img.size == (4, 4)
        assert img.mode == "RGBA"

    def test_is_transparent(self):
        data = _create_poster_image()
        img = Image.open(io.BytesIO(data))
        # All pixels should be fully transparent
        pixels = list(img.getdata())
        for pixel in pixels:
            assert pixel[3] == 0  # alpha channel


class TestConvertToPpsx:
    """Tests for _convert_to_ppsx."""

    def test_changes_content_type(self, tmp_path):
        # Create a minimal PPTX-like zip with [Content_Types].xml
        pptx_path = str(tmp_path / "test.pptx")
        ppsx_path = str(tmp_path / "test.ppsx")

        content_types_xml = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            b'<Override PartName="/ppt/presentation.xml" '
            b'ContentType="application/vnd.openxmlformats-officedocument.'
            b'presentationml.presentation.main+xml"/>'
            b'</Types>'
        )

        with zipfile.ZipFile(pptx_path, "w") as zf:
            zf.writestr("[Content_Types].xml", content_types_xml)
            zf.writestr("ppt/presentation.xml", b"<presentation/>")

        _convert_to_ppsx(pptx_path, ppsx_path)

        with zipfile.ZipFile(ppsx_path, "r") as zf:
            ct = zf.read("[Content_Types].xml")
            assert b"presentationml.slideshow.main+xml" in ct
            assert b"presentationml.presentation.main+xml" not in ct
            # Other files preserved
            assert zf.read("ppt/presentation.xml") == b"<presentation/>"

    def test_preserves_other_files(self, tmp_path):
        pptx_path = str(tmp_path / "test.pptx")
        ppsx_path = str(tmp_path / "test.ppsx")

        with zipfile.ZipFile(pptx_path, "w") as zf:
            zf.writestr("[Content_Types].xml", b"<Types/>")
            zf.writestr("ppt/slide1.xml", b"<slide>data</slide>")
            zf.writestr("ppt/media/image1.png", b"PNG_DATA")

        _convert_to_ppsx(pptx_path, ppsx_path)

        with zipfile.ZipFile(ppsx_path, "r") as zf:
            assert zf.read("ppt/slide1.xml") == b"<slide>data</slide>"
            assert zf.read("ppt/media/image1.png") == b"PNG_DATA"

    def test_with_real_pptx(self, tmp_path):
        """Test with a real python-pptx generated PPTX file."""
        pptx_path = tmp_path / "real.pptx"
        ppsx_path = tmp_path / "real.ppsx"
        _make_pptx(pptx_path, [{"text": "Test slide"}])

        _convert_to_ppsx(str(pptx_path), str(ppsx_path))

        with zipfile.ZipFile(str(ppsx_path), "r") as zf:
            ct = zf.read("[Content_Types].xml")
            assert b"presentationml.slideshow.main+xml" in ct
            assert b"presentationml.presentation.main+xml" not in ct


class TestSetSlideAdvanceTime:
    """Tests for _set_slide_advance_time."""

    def test_creates_correct_xml_new_transition(self, tmp_path):
        """When no transition element exists, one should be created."""
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"text": "Test"}])
        prs = Presentation(str(pptx_path))
        slide = prs.slides[0]

        _set_slide_advance_time(slide, 5000)

        p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        transitions = slide._element.findall(f"{{{p_ns}}}transition")
        assert len(transitions) == 1
        transition = transitions[0]
        assert transition.get("advTm") == "8000"  # 5000 + 3000 buffer

    def test_updates_existing_transition(self, tmp_path):
        """When a transition already exists, it should be updated."""
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"text": "Test"}])
        prs = Presentation(str(pptx_path))
        slide = prs.slides[0]

        # Set twice -- second call should update, not create duplicate
        _set_slide_advance_time(slide, 3000)
        _set_slide_advance_time(slide, 7000)

        p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        transitions = slide._element.findall(f"{{{p_ns}}}transition")
        assert len(transitions) == 1
        assert transitions[0].get("advTm") == "10000"  # 7000 + 3000

    def test_includes_3000ms_buffer(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"text": "Test"}])
        prs = Presentation(str(pptx_path))
        slide = prs.slides[0]

        _set_slide_advance_time(slide, 1000)

        p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        transition = slide._element.find(f"{{{p_ns}}}transition")
        assert transition.get("advTm") == "4000"  # 1000 + 3000

    def test_zero_duration(self, tmp_path):
        pptx_path = tmp_path / "test.pptx"
        _make_pptx(pptx_path, [{"text": "Test"}])
        prs = Presentation(str(pptx_path))
        slide = prs.slides[0]

        _set_slide_advance_time(slide, 0)

        p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        transition = slide._element.find(f"{{{p_ns}}}transition")
        assert transition.get("advTm") == "3000"  # 0 + 3000
