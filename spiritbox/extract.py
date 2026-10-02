"""Scope: extract — Extract slide PDF, text, and notes from PPTX."""

import shutil
import subprocess
from pathlib import Path

from pptx import Presentation

from .config import SpiritBoxConfig, get_source_pptx


def _extract_shape_text(shape) -> list[str]:
    """Extract text from a shape, including tables and groups."""
    parts = []
    if shape.has_text_frame:
        parts.append(shape.text_frame.text)
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in shape.shapes:
            parts.extend(_extract_shape_text(child))
    return parts


def _slide_title(slide, text_parts: list[str]) -> str:
    """Slide title: the title placeholder, else the first non-empty text line."""
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.has_text_frame:
        title = " ".join(title_shape.text_frame.text.split())
        if title:
            return title
    for part in text_parts:
        for line in part.splitlines():
            if line.strip():
                return " ".join(line.split())
    return ""


def extract_text_and_notes(pptx_path: str) -> list[dict]:
    """Extract text content and speaker notes from all slides."""
    prs = Presentation(pptx_path)
    slides_data = []

    for i, slide in enumerate(prs.slides, 1):
        text_parts = []
        for shape in slide.shapes:
            text_parts.extend(_extract_shape_text(shape))

        notes_text = ""
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text

        slides_data.append({
            "slide_number": i,
            "title": _slide_title(slide, text_parts),
            # Hidden slides are not exported to the PDF by LibreOffice.
            "hidden": slide._element.get("show") in ("0", "false"),
            "text": "\n".join(text_parts),
            "notes": notes_text,
        })

    return slides_data


def _find_libreoffice() -> str:
    """Find LibreOffice executable path."""
    mac_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if Path(mac_path).exists():
        return mac_path

    for cmd in ("soffice", "libreoffice"):
        found = shutil.which(cmd)
        if found:
            return found

    raise RuntimeError(
        "LibreOffice is required for PDF conversion.\n"
        "Install: brew install --cask libreoffice"
    )


def convert_to_pdf(pptx_path: str, output_dir: str) -> str:
    """Convert PPTX to PDF using LibreOffice.

    Gemini API accepts PDF directly for visual document understanding,
    so PNG conversion (and poppler dependency) is not needed.

    System deps: LibreOffice only (brew install --cask libreoffice)
    """
    soffice = _find_libreoffice()
    pptx_path = Path(pptx_path).resolve()
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf",
         "--outdir", str(output_dir), str(pptx_path)],
        check=True,
        capture_output=True,
        timeout=120,
    )

    pdf_path = output_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.exists():
        raise RuntimeError(f"PDF conversion failed: {pdf_path} not found")

    return str(pdf_path)


def extract_all(project_dir: str, config: SpiritBoxConfig = None) -> dict:
    """Extract all data from the source PPTX in the project directory. (P6: Resumable)

    Reads from source/ (the canonical copy created by init_project).

    Returns:
        {"slides_data": [...], "pdf_path": str, "slide_count": int}
    """
    pptx_path = get_source_pptx(project_dir, config)
    project = Path(project_dir)
    slides_dir = project / "slides"

    slides_data = extract_text_and_notes(str(pptx_path))
    slide_count = len(slides_data)

    # P6: Resumable — skip PDF conversion if already exists
    pdf_path = slides_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.exists():
        convert_to_pdf(str(pptx_path), str(slides_dir))

    return {
        "slides_data": slides_data,
        "pdf_path": str(pdf_path),
        "slide_count": slide_count,
    }
