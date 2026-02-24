"""Scope: assemble — Embed audio into PPTX and output as PPSX."""

import hashlib
import re
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.package import Part
from pptx.util import Cm

from .config import SpiritBoxConfig, get_source_pptx
from .voice import get_wav_duration_ms

# ---------------------------------------------------------------------------
# Monkey-patch: python-pptx 1.0.2 _find_by_sha1 crashes when iterating
# package media parts that lack a `sha1` property (base Part vs MediaPart).
# Adding a fallback sha1 property to the base Part class fixes it.
# ---------------------------------------------------------------------------
if not hasattr(Part, "sha1"):

    @property  # type: ignore[misc]
    def _part_sha1(self):
        return hashlib.sha1(self.blob).hexdigest()

    Part.sha1 = _part_sha1

# ---------------------------------------------------------------------------
# Register WAV audio content type (not included in python-pptx by default).
# Without this, add_movie(mime_type="audio/x-wav") raises KeyError.
# ---------------------------------------------------------------------------
if not hasattr(CT, "WAV"):
    CT.WAV = "audio/x-wav"

_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _create_poster_image() -> bytes:
    """Create a tiny transparent PNG for the audio shape poster frame."""
    img = Image.new("RGBA", (4, 4), (0, 0, 0, 0))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _get_shape_id(movie_shape) -> str | None:
    """Extract the shape ID from a movie shape element."""
    for el in movie_shape._element.iterdescendants():
        if etree.QName(el).localname == "cNvPr":
            return el.get("id")
    return None


def _set_slide_advance_time(slide, duration_ms: int):
    """Set slide auto-advance timing to match audio duration."""
    sld = slide._element
    p_ns = _P_NS

    existing = sld.findall(f"{{{p_ns}}}transition")
    if existing:
        transition = existing[0]
    else:
        transition = etree.SubElement(sld, f"{{{p_ns}}}transition")

        # Insert transition in correct schema order (after cSld)
        c_sld = sld.find(f"{{{p_ns}}}cSld")
        if c_sld is not None:
            insert_idx = list(sld).index(c_sld) + 1
            # Skip clrMapOvr if present
            clr = sld.find(f"{{{p_ns}}}clrMapOvr")
            if clr is not None:
                insert_idx = list(sld).index(clr) + 1
            sld.remove(transition)
            sld.insert(insert_idx, transition)

    # Audio start delay + audio duration + buffer (3000ms)
    transition.set("advTm", str(duration_ms + 3000))


def _strip_audio_from_source(pptx_path: str) -> str:
    """Remove all pre-existing audio from a PPTX at the ZIP level.

    Source PPTXs may contain audio from a previous SpiritBox run.
    This strips WAV media files, their relationship entries, audio shapes
    in slide XMLs, and timing nodes — preventing double-playback.

    Returns path to a cleaned temp copy (caller must delete).
    """
    clean_path = pptx_path + ".clean"

    with zipfile.ZipFile(pptx_path, "r") as zin:
        # Identify WAV media files
        wav_files = {
            f for f in zin.namelist()
            if re.match(r"ppt/media/.*\.wav$", f, re.IGNORECASE)
        }
        if not wav_files:
            return pptx_path  # Nothing to strip

        # Map WAV filenames to relationship targets for matching in .rels
        wav_targets = {f.split("/")[-1] for f in wav_files}

        with zipfile.ZipFile(clean_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                # Skip WAV media files entirely
                if item.filename in wav_files:
                    continue

                data = zin.read(item.filename)

                # Clean slide .rels: remove relationships pointing to WAVs
                if re.match(
                    r"ppt/slides/_rels/slide\d+\.xml\.rels$", item.filename
                ):
                    data = _clean_rels(data, wav_targets)

                # Clean slide XML: remove audio shapes and their timing
                if re.match(r"ppt/slides/slide\d+\.xml$", item.filename):
                    data = _clean_slide_xml(data)

                zout.writestr(item, data)

    return clean_path


def _clean_rels(data: bytes, wav_targets: set[str]) -> bytes:
    """Remove relationship entries that reference WAV files."""
    root = etree.fromstring(data)
    for rel in list(root):
        target = rel.get("Target", "")
        # Targets look like "../media/media1.wav"
        basename = target.split("/")[-1]
        if basename in wav_targets:
            root.remove(rel)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_slide_xml(data: bytes) -> bytes:
    """Remove audio-related shapes and timing from a slide XML."""
    root = etree.fromstring(data)
    nsmap = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }

    # Find shape tree
    sp_tree = root.find(".//p:cSld/p:spTree", nsmap)
    if sp_tree is not None:
        # Remove pic elements whose name ends with .wav
        for pic in list(sp_tree.findall(".//p:pic", nsmap)):
            nv_pr = pic.find(".//p:nvPicPr/p:cNvPr", nsmap)
            if nv_pr is None:
                nv_pr = pic.find(".//{%s}cNvPr" % nsmap["a"])
            if nv_pr is not None:
                name = nv_pr.get("name", "")
                if name.endswith(".wav"):
                    pic.getparent().remove(pic)
                    continue
            # Also check for audioFile elements inside the pic
            audio_file = pic.find(".//{%s}audioFile" % nsmap["a"])
            if audio_file is not None:
                pic.getparent().remove(pic)

    # Remove timing tree entirely — SpiritBox will recreate it
    timing = root.find("p:timing", nsmap)
    if timing is not None:
        root.remove(timing)

    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _embed_audio_in_slide(slide, audio_path: str, slide_width, slide_height) -> str | None:
    """Embed audio file in a slide. Returns shape ID for autoplay setup."""
    icon_size = Cm(1)
    left = slide_width - icon_size - Cm(0.5)
    top = slide_height - icon_size - Cm(0.5)

    # Create poster frame image (tiny transparent PNG)
    poster_path = Path(audio_path).parent / "_poster.png"
    if not poster_path.exists():
        poster_path.write_bytes(_create_poster_image())

    movie = slide.shapes.add_movie(
        audio_path,
        left, top, icon_size, icon_size,
        poster_frame_image=str(poster_path),
        mime_type="audio/x-wav",
    )

    return _get_shape_id(movie)


def _fix_autoplay_in_zip(pptx_path: str, audio_slide_nums: set[int]):
    """Post-process saved PPTX: change media delay from 'indefinite' to '0'.

    python-pptx's add_movie creates timing with delay="indefinite" (click to play).
    Changing it to delay="0" triggers playback on slide entry.
    This is the minimal change — keeps the rest of python-pptx's XML intact.
    """
    tmp_path = pptx_path + ".tmp"

    with zipfile.ZipFile(pptx_path, "r") as zin:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)

                match = re.match(r"ppt/slides/slide(\d+)\.xml", item.filename)
                if match:
                    slide_num = int(match.group(1))
                    if slide_num in audio_slide_nums:
                        # Only change: delay="indefinite" → delay="0"
                        data = data.replace(
                            b'delay="indefinite"',
                            b'delay="0"',
                        )

                zout.writestr(item, data)

    Path(tmp_path).replace(pptx_path)


def _convert_to_ppsx(pptx_path: str, ppsx_path: str):
    """Convert PPTX to PPSX by changing the content type."""
    with zipfile.ZipFile(pptx_path, "r") as zin:
        with zipfile.ZipFile(ppsx_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"presentationml.presentation.main+xml",
                        b"presentationml.slideshow.main+xml",
                    )
                zout.writestr(item, data)


def assemble_ppsx(
    project_dir: str,
    durations: dict[int, int],
    config: SpiritBoxConfig = None,
) -> str:
    """Embed audio into PPTX and output as PPSX. (P8: Non-Destructive)

    Args:
        project_dir: project directory
        durations: {slide_number: duration_ms} from voice generation
        config: SpiritBox config (for resolving source PPTX)

    Returns: path to the output PPSX file
    """
    project = Path(project_dir)
    audio_dir = project / "audio"
    output_dir = project / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    pptx_path = get_source_pptx(project_dir, config)

    # Strip any pre-existing audio from source to prevent double-playback
    clean_path = _strip_audio_from_source(str(pptx_path))
    try:
        prs = Presentation(clean_path)
        slide_width = prs.slide_width
        slide_height = prs.slide_height

        audio_slide_nums = set()

        for i, slide in enumerate(prs.slides, 1):
            wav_path = audio_dir / f"slide_{i:03d}.wav"
            if not wav_path.exists():
                _set_slide_advance_time(slide, 3000)
                continue

            duration_ms = durations.get(i, get_wav_duration_ms(str(wav_path)))

            shape_id = _embed_audio_in_slide(
                slide, str(wav_path), slide_width, slide_height
            )
            if shape_id:
                audio_slide_nums.add(i)

            _set_slide_advance_time(slide, duration_ms)

        # Save as PPTX (python-pptx creates click-triggered timing)
        tmp_pptx = output_dir / f"{pptx_path.stem}_narrated.pptx"
        prs.save(str(tmp_pptx))

        # Minimal post-process: change delay="indefinite" to delay="0"
        _fix_autoplay_in_zip(str(tmp_pptx), audio_slide_nums)

        # Convert to PPSX (avoid overwriting existing files)
        base_name = f"{pptx_path.stem}_narrated"
        ppsx_path = output_dir / f"{base_name}.ppsx"
        n = 2
        while ppsx_path.exists():
            ppsx_path = output_dir / f"{base_name}_{n}.ppsx"
            n += 1

        try:
            _convert_to_ppsx(str(tmp_pptx), str(ppsx_path))
        finally:
            if tmp_pptx.exists():
                tmp_pptx.unlink()

    finally:
        # Clean up temp file if we created one
        if clean_path != str(pptx_path) and Path(clean_path).exists():
            Path(clean_path).unlink()

    return str(ppsx_path)
